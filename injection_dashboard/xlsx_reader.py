from __future__ import annotations

import io, re
from typing import Any, Optional, Tuple, Union
from openpyxl import load_workbook

MONTH_RE = re.compile(r"(?P<year>20\d{2})\s*년?\s*(?P<month>1[0-2]|[1-9])\s*월")


def _wb(source: Union[bytes, str]):
    if isinstance(source, (bytes, bytearray)):
        return load_workbook(io.BytesIO(bytes(source)), data_only=True, read_only=True)
    return load_workbook(source, data_only=True, read_only=True)


def _f(v):
    if v in (None, ""): return None
    try: return float(v)
    except: return None


def _i(v):
    x=_f(v); return int(round(x)) if x is not None else None


def parse_year_month(text: Any, fallback_year: Optional[int]=None) -> Tuple[Optional[int], Optional[int]]:
    if text is None: return fallback_year, None
    s=str(text).strip(); m=MONTH_RE.search(s)
    if m: return int(m.group('year')), int(m.group('month'))
    m=re.search(r"(1[0-2]|[1-9])\s*월",s)
    return (fallback_year, int(m.group(1))) if m else (fallback_year,None)


def _latest_month(ws, rows, first_col):
    latest=None
    for m in range(1,13):
        c=first_col+m-1
        if any(_f(ws.cell(r,c).value) is not None for r in rows): latest=m
    return latest


def parse_labor(source):
    wb=_wb(source)
    ws=wb['공정요약(월별)']
    _,m1=parse_year_month(ws['B2'].value)
    y,m2=parse_year_month(ws['D2'].value)
    month=m2 or m1
    if y is None or month is None: raise ValueError('노무보고서 대상 연월 확인 실패')
    g=lambda c:_f(ws[c].value)
    metrics={
      'headcount':_i(ws['B7'].value),'workdays':_i(ws['B24'].value),'nominal_hours':g('B25'),
      'available_hours':g('B32'),'vacation_hours':abs(g('B28') or 0),'support_out_hours':abs(g('B29') or 0),
      'support_in_hours':g('B30') or 0,'overtime_hours':g('B31') or 0,'capacity_loss_rate':g('B34'),
      'erp_standard_hours':g('B37'),'erp_actual_hours':g('B38'),'non_erp_hours':g('B39'),'operating_hours':g('B40'),
      'unclassified_hours':g('B41'),'direct_work_rate':g('B44'),'time_allocation_rate':g('B45'),
      'capacity_utilization_rate':g('B46'),'work_efficiency':g('B47'),'unclassified_rate':g('B48'),
      'manufacturing_equipment_rate':g('B18'),'packaging_equipment_rate':g('B19')}
    processes=[]
    for r in range(53,59):
        vals=[ws.cell(r,c).value for c in range(1,8)]
        if not vals[1]: continue
        processes.append({'part':vals[0],'process':vals[1],'standard_hours':_f(vals[2]),'actual_hours':_f(vals[3]),'delta_hours':_f(vals[4]),'efficiency':_f(vals[5]),'share':_f(vals[6])})
    history=[]
    if '공정요약(기간)' in wb.sheetnames:
        wp=wb['공정요약(기간)']
        for r in range(6,18):
            _,mm=parse_year_month(wp[f'J{r}'].value,y)
            if mm is None: continue
            available=_f(wp[f'R{r}'].value); standard=_f(wp[f'S{r}'].value); actual=_f(wp[f'T{r}'].value); non=_f(wp[f'U{r}'].value); nominal=_f(wp[f'M{r}'].value)
            if not available: continue
            op=(actual or 0)+(non or 0)
            history.append({'month':mm,'work_efficiency':standard/actual*100 if standard and actual else None,'time_allocation_rate':op/available*100,'capacity_utilization_rate':op/nominal*100 if nominal else None,'available_hours':available,'erp_actual_hours':actual,'non_erp_hours':non,'unclassified_hours':available-op})
    return {'year':y,'month':month,'metrics':metrics,'processes':processes,'history':history}


def _series(ws,row,first_col): return [_f(ws.cell(row,first_col+m-1).value) for m in range(1,13)]


def parse_yield(source,fallback_year=None):
    wb=_wb(source); wi=wb['2_이물검사_수율']; wp=wb['3_포장_수율']; wd=wb['4_종합_대시보드']; wm=wb['5_월간_수율현황']
    month=max(x for x in [_latest_month(wi,[12,13,14],2),_latest_month(wp,[12,13,14],2)] if x is not None)
    specs=['12mL','15mL','20mL']; rr={'12mL':12,'15mL':13,'20mL':14}; sm={'12mL':6,'15mL':7,'20mL':8}; col=1+month
    current={s:{'inspection_yield':_f(wi.cell(rr[s],col).value),'packaging_yield':_f(wp.cell(rr[s],col).value)} for s in specs}
    cumulative={}
    for s in specs:
        r=sm[s]
        cumulative[s]={'inspection_batches':_i(wi[f'B{r}'].value),'inspection_avg_yield':_f(wi[f'C{r}'].value),'inspection_limit':_f(wi[f'F{r}'].value),'inspection_deviations':_i(wi[f'G{r}'].value),'inspection_loss':_i(wi[f'J{r}'].value),'packaging_batches':_i(wp[f'B{r}'].value),'packaging_avg_yield':_f(wp[f'C{r}'].value),'packaging_limit':_f(wp[f'F{r}'].value),'packaging_deviations':_i(wp[f'G{r}'].value),'packaging_loss':_i(wp[f'K{r}'].value)}
    total={'batches':_i(wm['B15'].value),'ordered_qty':_i(wm['C15'].value),'actual_qty':_i(wm['D15'].value),'packaging_yield':_f(wm['E15'].value),'deviation_batches':_i(wm['I15'].value),'packaging_loss':_i(wm['J15'].value)}
    dev=sum(_i(wd.cell(r,col).value) or 0 for r in (14,15,16))
    return {'year':fallback_year,'month':month,'current':current,'cumulative':cumulative,'total':total,'current_deviation_count':dev,'inspection_trend':{s:_series(wi,rr[s],2) for s in specs},'packaging_trend':{s:_series(wp,rr[s],2) for s in specs}}


def parse_mbo(source,fallback_year=None):
    wb=_wb(source); ws=wb['2_성과요약']; wm=wb['3_월별모니터링']; month=_latest_month(wm,[6,9,12],3); col=2+month
    rows={'12mL':(6,8),'15mL':(9,11),'20mL':(12,14)}
    current={s:{'hours_per_batch':_f(wm.cell(a,col).value),'reduction_rate':_f(wm.cell(b,col).value)} for s,(a,b) in rows.items()}
    annual={}
    for s,r in {'12mL':12,'15mL':13,'20mL':14}.items():
        annual[s]={'baseline_2025_hours':_f(ws[f'C{r}'].value),'current_2026_hours':_f(ws[f'E{r}'].value),'reduction_rate':_f(ws[f'G{r}'].value),'target_hours':_f(ws[f'H{r}'].value),'status':ws[f'I{r}'].value}
    overall={'baseline_2025_hours':_f(ws['C15'].value),'current_2026_hours':_f(ws['E15'].value),'saved_hours':_f(ws['F15'].value),'reduction_rate':_f(ws['G15'].value),'target_reduction_rate':30.0,'status':ws['D19'].value}
    trend={'12mL':_series(wm,6,3),'15mL':_series(wm,9,3),'20mL':_series(wm,12,3)}
    targets={'12mL':_f(wm['C7'].value),'15mL':_f(wm['C10'].value),'20mL':_f(wm['C13'].value)}
    return {'year':fallback_year,'month':month,'current':current,'annual':annual,'overall':overall,'trend':trend,'targets':targets}


def parse_bundle(labor_source,yield_source,mbo_source):
    labor=parse_labor(labor_source); yy=parse_yield(yield_source,labor['year']); mbo=parse_mbo(mbo_source,labor['year'])
    months={'노무':labor['month'],'수율':yy['month'],'이물검사':mbo['month']}
    if len(set(months.values()))!=1: raise ValueError('3개 파일의 대상월이 일치하지 않습니다: '+', '.join(f'{k}={v}월' for k,v in months.items()))
    return {'year':labor['year'],'month':labor['month'],'labor':labor,'yield':yy,'mbo':mbo}
